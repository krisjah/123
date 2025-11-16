#!/usr/bin/env python3
"""
Beer Brewery Financial Model - Operations & Costs
Builds: Production_Capacity, COGS_Buildup, Logistics_Distribution, SG&A sheets

Run after: 02_volume_pricing.py
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Color definitions
BLUE_INPUT = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
GREY_CALC = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
YELLOW_OUTPUT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
DARK_BLUE_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

MONTHS = 120
TIMELINE_START_COL = 6


def add_timeline_header(ws, start_col=6, num_months=120):
    """Add standard timeline header rows"""
    ws['A3'] = 'Period'
    ws['A4'] = 'Date'
    ws['A5'] = 'Year'
    ws['A6'] = 'FY'
    ws['A7'] = 'Quarter'

    for month in range(1, num_months + 1):
        col = start_col + month - 1
        col_letter = get_column_letter(col)
        ws[f'{col_letter}3'] = month
        ws[f'{col_letter}3'].font = Font(size=9)
        ws[f'{col_letter}4'] = f'=EDATE(Control!$B$4,{month-1})'
        ws[f'{col_letter}4'].number_format = 'mmm-yy'
        ws[f'{col_letter}4'].font = Font(bold=True)
        ws[f'{col_letter}5'] = f'=({month}-1)/12'
        ws[f'{col_letter}5'].number_format = '0.00'
        ws[f'{col_letter}5'].font = Font(size=9, color="808080")
        ws[f'{col_letter}6'] = f'=YEAR({col_letter}4)'
        ws[f'{col_letter}6'].font = Font(bold=True)
        ws[f'{col_letter}7'] = f'="Q"&ROUNDUP(MONTH({col_letter}4)/3,0)'
        ws[f'{col_letter}7'].font = Font(bold=True)


def build_production_capacity_sheet(wb):
    """Build Production_Capacity sheet"""

    print("Building Production_Capacity sheet...")
    ws = wb['Production_Capacity']

    # Title
    ws['A1'] = 'PRODUCTION CAPACITY'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # Capacity assumptions
    ws['A10'] = 'Brewery Capacity'
    ws['A10'].font = Font(bold=True, size=11)

    headers = ['', 'Base', 'Upside', 'Downside', 'Active']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    ws['A12'] = 'Installed brewhouse capacity (HL/year)'
    ws['B12'] = 1000000
    ws['C12'] = 1000000
    ws['D12'] = 1000000
    ws['E12'] = '=INDEX(B12:D12,1,Control!$B$10)'

    for col in ['B', 'C', 'D']:
        ws[f'{col}12'].number_format = '#,##0'
        ws[f'{col}12'].fill = BLUE_INPUT
    ws['E12'].number_format = '#,##0'
    ws['E12'].fill = GREY_CALC

    ws['A13'] = 'Installed capacity (HL/month)'
    ws['B13'] = '=B12/12'
    ws['C13'] = '=C12/12'
    ws['D13'] = '=D12/12'
    ws['E13'] = '=E12/12'

    for col in ['B', 'C', 'D', 'E']:
        ws[f'{col}13'].number_format = '#,##0'

    # Yield & Loss Factors
    ws['A15'] = 'Yield & Loss Factors'
    ws['A15'].font = Font(bold=True, size=11)

    yield_data = [
        ('Brewhouse yield %', 0.95, 0.96, 0.94),
        ('Fermentation yield %', 0.98, 0.98, 0.97),
        ('Packaging yield %', 0.97, 0.98, 0.96),
    ]

    row = 16
    for label, base, upside, downside in yield_data:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = base
        ws[f'C{row}'] = upside
        ws[f'D{row}'] = downside
        ws[f'E{row}'] = f'=INDEX(B{row}:D{row},1,Control!$B$10)'

        for col in ['B', 'C', 'D', 'E']:
            ws[f'{col}{row}'].number_format = '0.0%'
        for col in ['B', 'C', 'D']:
            ws[f'{col}{row}'].fill = BLUE_INPUT
        ws[f'E{row}'].fill = GREY_CALC
        row += 1

    ws['A19'] = 'Cumulative yield %'
    ws['E19'] = '=E16*E17*E18'
    ws['E19'].number_format = '0.0%'
    ws['E19'].font = Font(bold=True)

    # Monthly calculations
    ws['A22'] = 'Monthly Production & Utilization'
    ws['A22'].font = Font(bold=True, size=11)

    calc_labels = [
        'HL Sold (from Volume sheet)',
        'HL to be Brewed (grossed up for losses)',
        'Installed capacity (HL/month)',
        'Capacity utilization %',
        'Headroom (HL)',
        'Constraint flag'
    ]

    for idx, label in enumerate(calc_labels, start=23):
        ws[f'A{idx}'] = label

    # Formulas for each month
    for month in range(1, MONTHS + 1):
        col = TIMELINE_START_COL + month - 1
        col_letter = get_column_letter(col)

        # HL Sold
        ws[f'{col_letter}23'] = f'=Volume_Assumptions!{col_letter}51'
        ws[f'{col_letter}23'].number_format = '#,##0'

        # HL Brewed
        ws[f'{col_letter}24'] = f'={col_letter}23/$E$19'
        ws[f'{col_letter}24'].number_format = '#,##0'

        # Capacity
        ws[f'{col_letter}25'] = '=$E$13'
        ws[f'{col_letter}25'].number_format = '#,##0'

        # Utilization %
        ws[f'{col_letter}26'] = f'={col_letter}24/{col_letter}25'
        ws[f'{col_letter}26'].number_format = '0.0%'

        # Headroom
        ws[f'{col_letter}27'] = f'={col_letter}25-{col_letter}24'
        ws[f'{col_letter}27'].number_format = '#,##0'

        # Flag
        ws[f'{col_letter}28'] = f'=IF({col_letter}26>1,"OVER CAPACITY","OK")'

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18

    print("  ✓ Production_Capacity sheet complete")


def build_cogs_buildup_sheet(wb):
    """Build COGS_Buildup sheet"""

    print("Building COGS_Buildup sheet...")
    ws = wb['COGS_Buildup']

    # Title
    ws['A1'] = 'COGS BUILDUP'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # Materials cost assumptions
    ws['A10'] = 'Materials Cost ($/HL by brand)'
    ws['A10'].font = Font(bold=True, size=11)

    headers = ['Brand', 'Year 1 $/HL', 'Infl % Base', 'Upside', 'Downside']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    brands_materials = [
        ('SkyBrew', 25.00, 0.025, 0.035, 0.020),
        ('AllDark', 32.00, 0.025, 0.035, 0.020),
        ('GoldCoast', 28.00, 0.025, 0.035, 0.020),
        ('CityLight', 22.00, 0.025, 0.035, 0.020),
        ('HarborIPA', 35.00, 0.025, 0.035, 0.020),
    ]

    row = 12
    for brand, y1_cost, base_infl, up_infl, down_infl in brands_materials:
        ws[f'A{row}'] = brand
        ws[f'B{row}'] = y1_cost
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'B{row}'].fill = BLUE_INPUT

        ws[f'C{row}'] = base_infl
        ws[f'D{row}'] = up_infl
        ws[f'E{row}'] = down_infl

        for col in ['C', 'D', 'E']:
            ws[f'{col}{row}'].number_format = '0.0%'
            ws[f'{col}{row}'].fill = BLUE_INPUT
        row += 1

    # Packaging cost assumptions
    pkg_row = row + 1
    ws[f'A{pkg_row}'] = 'Packaging Cost ($/HL by pack type)'
    ws[f'A{pkg_row}'].font = Font(bold=True, size=11)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=pkg_row+1, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    packs_costs = [
        ('Bottles', 18.00, 0.030, 0.040, 0.025),
        ('Cans', 15.00, 0.030, 0.040, 0.025),
        ('Kegs', 8.00, 0.020, 0.030, 0.015),
    ]

    row = pkg_row + 2
    for pack, y1_cost, base_infl, up_infl, down_infl in packs_costs:
        ws[f'A{row}'] = pack
        ws[f'B{row}'] = y1_cost
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'B{row}'].fill = BLUE_INPUT

        ws[f'C{row}'] = base_infl
        ws[f'D{row}'] = up_infl
        ws[f'E{row}'] = down_infl

        for col in ['C', 'D', 'E']:
            ws[f'{col}{row}'].number_format = '0.0%'
            ws[f'{col}{row}'].fill = BLUE_INPUT
        row += 1

    # Labour costs
    labour_row = row + 1
    ws[f'A{labour_row}'] = 'Direct Labour'
    ws[f'A{labour_row}'].font = Font(bold=True, size=11)

    ws[f'A{labour_row+1}'] = 'Headcount'
    ws[f'B{labour_row+1}'] = 50
    ws[f'C{labour_row+1}'] = 52
    ws[f'D{labour_row+1}'] = 48
    ws[f'E{labour_row+1}'] = '=INDEX(B{0}:D{0},1,Control!$B$10)'.format(labour_row+1)

    for col in ['B', 'C', 'D']:
        ws[f'{col}{labour_row+1}'].fill = BLUE_INPUT
    ws[f'E{labour_row+1}'].fill = GREY_CALC

    ws[f'A{labour_row+2}'] = 'Avg wage ($/year)'
    ws[f'B{labour_row+2}'] = 60000
    ws[f'C{labour_row+2}'] = 60000
    ws[f'D{labour_row+2}'] = 60000
    ws[f'E{labour_row+2}'] = '=INDEX(B{0}:D{0},1,Control!$B$10)'.format(labour_row+2)

    for col in ['B', 'C', 'D']:
        ws[f'{col}{labour_row+2}'].number_format = '$#,##0'
        ws[f'{col}{labour_row+2}'].fill = BLUE_INPUT
    ws[f'E{labour_row+2}'].number_format = '$#,##0'
    ws[f'E{labour_row+2}'].fill = GREY_CALC

    # Utilities
    util_row = labour_row + 4
    ws[f'A{util_row}'] = 'Utilities'
    ws[f'A{util_row}'].font = Font(bold=True, size=11)

    ws[f'A{util_row+1}'] = 'Utilities cost per HL brewed ($/HL)'
    ws[f'B{util_row+1}'] = 3.00
    ws[f'C{util_row+1}'] = 3.00
    ws[f'D{util_row+1}'] = 3.00

    for col in ['B', 'C', 'D']:
        ws[f'{col}{util_row+1}'].number_format = '$#,##0.00'
        ws[f'{col}{util_row+1}'].fill = BLUE_INPUT

    # Overhead
    oh_row = util_row + 3
    ws[f'A{oh_row}'] = 'Manufacturing Overhead'
    ws[f'A{oh_row}'].font = Font(bold=True, size=11)

    ws[f'A{oh_row+1}'] = 'Fixed overhead ($/month)'
    ws[f'B{oh_row+1}'] = 200000
    ws[f'C{oh_row+1}'] = 200000
    ws[f'D{oh_row+1}'] = 200000

    for col in ['B', 'C', 'D']:
        ws[f'{col}{oh_row+1}'].number_format = '$#,##0'
        ws[f'{col}{oh_row+1}'].fill = BLUE_INPUT

    ws[f'A{oh_row+2}'] = 'Variable overhead ($/HL)'
    ws[f'B{oh_row+2}'] = 2.00
    ws[f'C{oh_row+2}'] = 2.00
    ws[f'D{oh_row+2}'] = 2.00

    for col in ['B', 'C', 'D']:
        ws[f'{col}{oh_row+2}'].number_format = '$#,##0.00'
        ws[f'{col}{oh_row+2}'].fill = BLUE_INPUT

    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 18

    print("  ✓ COGS_Buildup sheet complete")
    print("    (Note: Detailed COGS calculations by SKU will be in future version)")


def build_logistics_distribution_sheet(wb):
    """Build Logistics_Distribution sheet"""

    print("Building Logistics_Distribution sheet...")
    ws = wb['Logistics_Distribution']

    # Title
    ws['A1'] = 'LOGISTICS & DISTRIBUTION'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # Transport cost assumptions
    ws['A10'] = 'Transport Cost ($/HL)'
    ws['A10'].font = Font(bold=True, size=11)

    headers = ['Channel', 'Year 1 $/HL', 'Infl % (all scenarios)']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    ws['A12'] = 'OnTrade'
    ws['B12'] = 8.00
    ws['C12'] = 0.03

    ws['A13'] = 'OffTrade'
    ws['B13'] = 5.00
    ws['C13'] = 0.03

    for row in [12, 13]:
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'B{row}'].fill = BLUE_INPUT
        ws[f'C{row}'].number_format = '0.0%'
        ws[f'C{row}'].fill = BLUE_INPUT

    # Warehouse costs
    ws['A15'] = 'Warehouse Costs'
    ws['A15'].font = Font(bold=True, size=11)

    ws['A16'] = 'Fixed rental ($/month)'
    ws['B16'] = 50000
    ws['C16'] = 0.025

    ws['A17'] = 'Variable handling ($/HL)'
    ws['B17'] = 1.00
    ws['C17'] = 0.025

    for row in [16, 17]:
        ws[f'B{row}'].number_format = '$#,##0.00'
        ws[f'B{row}'].fill = BLUE_INPUT
        ws[f'C{row}'].number_format = '0.0%'
        ws[f'C{row}'].fill = BLUE_INPUT

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 20

    print("  ✓ Logistics_Distribution sheet complete")


def build_sga_sheet(wb):
    """Build SG&A sheet"""

    print("Building SG&A sheet...")
    ws = wb['SG&A']

    # Title
    ws['A1'] = 'SG&A (SALES, GENERAL & ADMIN)'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # A&P assumptions
    ws['A10'] = 'A&P as % of Net Revenue'
    ws['A10'].font = Font(bold=True, size=11)

    headers = ['Period', 'Base', 'Upside', 'Downside']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    ws['A12'] = 'Year 1-3'
    ws['B12'] = 0.08
    ws['C12'] = 0.07
    ws['D12'] = 0.09

    ws['A13'] = 'Year 4+'
    ws['B13'] = 0.07
    ws['C13'] = 0.06
    ws['D13'] = 0.08

    for row in [12, 13]:
        for col in ['B', 'C', 'D']:
            ws[f'{col}{row}'].number_format = '0.0%'
            ws[f'{col}{row}'].fill = BLUE_INPUT

    # Headcount - Sales
    ws['A15'] = 'Sales Team'
    ws['A15'].font = Font(bold=True, size=11)

    ws['A16'] = 'Headcount'
    ws['B16'] = 15
    ws['C16'] = 16
    ws['D16'] = 14

    ws['A17'] = 'Avg salary ($/year)'
    ws['B17'] = 70000
    ws['C17'] = 70000
    ws['D17'] = 70000

    for row in [16, 17]:
        for col in ['B', 'C', 'D']:
            ws[f'{col}{row}'].fill = BLUE_INPUT
            if row == 17:
                ws[f'{col}{row}'].number_format = '$#,##0'

    # Headcount - Admin
    ws['A19'] = 'Admin Team'
    ws['A19'].font = Font(bold=True, size=11)

    ws['A20'] = 'Headcount'
    ws['B20'] = 20
    ws['C20'] = 22
    ws['D20'] = 18

    ws['A21'] = 'Avg salary ($/year)'
    ws['B21'] = 80000
    ws['C21'] = 80000
    ws['D21'] = 80000

    for row in [20, 21]:
        for col in ['B', 'C', 'D']:
            ws[f'{col}{row}'].fill = BLUE_INPUT
            if row == 21:
                ws[f'{col}{row}'].number_format = '$#,##0'

    # Office costs
    ws['A23'] = 'Office & Other'
    ws['A23'].font = Font(bold=True, size=11)

    ws['A24'] = 'Monthly cost'
    ws['B24'] = 30000
    ws['C24'] = 30000
    ws['D24'] = 30000

    for col in ['B', 'C', 'D']:
        ws[f'{col}24'].number_format = '$#,##0'
        ws[f'{col}24'].fill = BLUE_INPUT

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15

    print("  ✓ SG&A sheet complete")


def main():
    """Main execution"""
    print("\n" + "="*60)
    print("Building Operations & Costs Sheets")
    print("="*60 + "\n")

    wb = load_workbook('Brewery_Financial_Model.xlsx')

    build_production_capacity_sheet(wb)
    build_cogs_buildup_sheet(wb)
    build_logistics_distribution_sheet(wb)
    build_sga_sheet(wb)

    wb.save('Brewery_Financial_Model.xlsx')

    print("\n✅ Operations & Costs sheets complete!")
    print("\n   - Production capacity tracking")
    print("   - COGS structure (materials, packaging, labour, overhead)")
    print("   - Logistics & distribution costs")
    print("   - SG&A (A&P, headcount, office)")
    print("\nNext step: Run 04_capital_funding.py")


if __name__ == '__main__':
    main()
