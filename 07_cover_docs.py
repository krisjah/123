#!/usr/bin/env python3
"""
Beer Brewery Financial Model - Cover & Documentation
Builds: Cover, Documentation, Change_Log sheets

Run after: 06_kpis_checks.py
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date

DARK_BLUE_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
LIGHT_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")


def build_cover_sheet(wb):
    """Build Cover sheet"""

    print("Building Cover sheet...")
    ws = wb['Cover']

    # Title
    ws['A1'] = 'BEER BREWERY FINANCIAL MODEL'
    ws['A1'].font = Font(bold=True, size=20, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells('A1:E1')
    ws.row_dimensions[1].height = 30

    ws['A2'] = 'Australia & New Zealand Operations'
    ws['A2'].font = Font(size=14, color="4472C4")
    ws['A2'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A2:E2')

    # Model info
    ws['A4'] = 'Model Version:'
    ws['B4'] = '1.0'
    ws['B4'].font = Font(bold=True)

    ws['A5'] = 'Date Created:'
    ws['B5'] = date.today().strftime('%d-%b-%Y')

    ws['A6'] = 'Currency:'
    ws['B6'] = 'AUD'

    ws['A7'] = 'Projection Period:'
    ws['B7'] = '10 years (120 months)'

    # Navigation
    ws['A10'] = 'NAVIGATION'
    ws['A10'].font = Font(bold=True, size=12, color="FFFFFF")
    ws['A10'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A10:C10')

    nav_items = [
        ('Control & Inputs', 'Control'),
        ('Assumptions (Global)', 'Assumptions_Global'),
        ('Product Matrix', 'Product_Matrix'),
        ('Volume Assumptions', 'Volume_Assumptions'),
        ('Pricing & Revenue', 'Pricing_Revenue'),
        ('Production & Capacity', 'Production_Capacity'),
        ('COGS', 'COGS_Buildup'),
        ('Logistics', 'Logistics_Distribution'),
        ('SG&A', 'SG&A'),
        ('Capex', 'Capex_Schedule'),
        ('Working Capital', 'Working_Capital'),
        ('Debt', 'Debt_Schedule'),
        ('Equity', 'Equity'),
        ('P&L', 'P&L'),
        ('Cash Flow', 'CashFlow'),
        ('Balance Sheet', 'BalanceSheet'),
        ('KPI Dashboard', 'KPI_Dashboard'),
        ('Checks & Errors', 'Checks_Errors'),
    ]

    row = 11
    for label, sheet_name in nav_items:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = '→'
        ws[f'C{row}'] = f'=HYPERLINK("#\'{sheet_name}\'!A1","Go to {sheet_name}")'
        ws[f'C{row}'].font = Font(color="0563C1", underline="single")
        ws[f'C{row}'].fill = LIGHT_GREEN
        row += 1

    # Instructions
    ws[f'A{row+2}'] = 'USER INSTRUCTIONS'
    ws[f'A{row+2}'].font = Font(bold=True, size=12, color="FFFFFF")
    ws[f'A{row+2}'].fill = DARK_BLUE_HEADER
    ws.merge_cells(f'A{row+2}:E{row+2}')

    instructions = [
        '1. Start with CONTROL sheet - select scenario (Base/Upside/Downside)',
        '2. Review all BLUE cells (inputs) in assumption sheets',
        '3. All other cells are calculated - do not edit',
        '4. Enable iterative calculation: File → Options → Formulas',
        '5. Check Checks_Errors sheet for model integrity',
        '6. Review outputs in P&L, CashFlow, BalanceSheet, KPIs',
    ]

    row += 3
    for instruction in instructions:
        ws[f'A{row}'] = instruction
        ws[f'A{row}'].alignment = Alignment(wrap_text=True)
        row += 1

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 5
    ws.column_dimensions['C'].width = 30

    print("  ✓ Cover sheet complete")


def build_documentation_sheet(wb):
    """Build Documentation sheet"""

    print("Building Documentation sheet...")
    ws = wb['Documentation']

    # Title
    ws['A1'] = 'DOCUMENTATION'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:C1')

    # Model overview
    ws['A3'] = 'Model Overview'
    ws['A3'].font = Font(bold=True, size=12)

    overview = [
        ('Purpose', 'Financial model for beer brewery and distribution business in Australia/New Zealand'),
        ('Users', 'Equity investors, lenders, management, board'),
        ('Projection period', '10 years (120 months)'),
        ('Currency', 'AUD'),
        ('Accounting standard', 'IFRS'),
    ]

    row = 4
    for label, value in overview:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

    # Key assumptions
    ws[f'A{row+1}'] = 'Key Assumptions (Base Scenario)'
    ws[f'A{row+1}'].font = Font(bold=True, size=12)

    assumptions = [
        ('Volume growth', '3% p.a. compounded'),
        ('Price escalation', '2.5% p.a.'),
        ('Materials inflation', '2.5% p.a.'),
        ('Wage inflation', '3% p.a.'),
        ('Corporate tax rate', '30%'),
        ('Excise duty', '$58.26 per litre of alcohol'),
        ('Discount rate (WACC)', '12%'),
    ]

    row += 2
    for label, value in assumptions:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 50

    print("  ✓ Documentation sheet complete")


def build_changelog_sheet(wb):
    """Build Change_Log sheet"""

    print("Building Change_Log sheet...")
    ws = wb['Change_Log']

    # Title
    ws['A1'] = 'CHANGE LOG'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    # Headers
    headers = ['Date', 'Version', 'Author', 'Change Description', 'Sheets Affected']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    # Initial entry
    ws['A4'] = date.today().strftime('%d-%b-%Y')
    ws['B4'] = '1.0'
    ws['C4'] = 'Model Builder'
    ws['D4'] = 'Initial model build - Python-generated'
    ws['E4'] = 'All'

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 20

    print("  ✓ Change_Log sheet complete")


def main():
    """Main execution"""
    print("\n" + "="*60)
    print("Building Cover & Documentation Sheets")
    print("="*60 + "\n")

    wb = load_workbook('Brewery_Financial_Model.xlsx')

    build_cover_sheet(wb)
    build_documentation_sheet(wb)
    build_changelog_sheet(wb)

    wb.save('Brewery_Financial_Model.xlsx')

    print("\n✅ Cover & Documentation sheets complete!")
    print("\n   - Cover page with navigation")
    print("   - Documentation of assumptions")
    print("   - Change log")
    print("\nNext step: Run 08_formatting.py")


if __name__ == '__main__':
    main()
