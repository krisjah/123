#!/usr/bin/env python3
"""
Beer Brewery Financial Model - Volume & Pricing
Builds: Volume_Assumptions, Pricing_Revenue sheets

Run after: 01_control_assumptions.py
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Color definitions
BLUE_INPUT = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
GREY_CALC = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
YELLOW_OUTPUT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
DARK_BLUE_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

# 120 months
MONTHS = 120
# Timeline starts at column F
TIMELINE_START_COL = 6  # Column F


def add_timeline_header(ws, start_col=6, num_months=120):
    """Add standard timeline header rows (rows 3-7)"""

    ws['A3'] = 'Period'
    ws['A4'] = 'Date'
    ws['A5'] = 'Year'
    ws['A6'] = 'FY'
    ws['A7'] = 'Quarter'

    for month in range(1, num_months + 1):
        col = start_col + month - 1
        col_letter = get_column_letter(col)

        # Period number
        ws[f'{col_letter}3'] = month
        ws[f'{col_letter}3'].font = Font(size=9)

        # Date
        ws[f'{col_letter}4'] = f'=EDATE(Control!$B$4,{month-1})'
        ws[f'{col_letter}4'].number_format = 'mmm-yy'
        ws[f'{col_letter}4'].font = Font(bold=True)

        # Year (decimal)
        ws[f'{col_letter}5'] = f'=({month}-1)/12'
        ws[f'{col_letter}5'].number_format = '0.00'
        ws[f'{col_letter}5'].font = Font(size=9, color="808080")

        # Fiscal year
        ws[f'{col_letter}6'] = f'=YEAR({col_letter}4)'
        ws[f'{col_letter}6'].font = Font(bold=True)

        # Quarter
        ws[f'{col_letter}7'] = f'="Q"&ROUNDUP(MONTH({col_letter}4)/3,0)'
        ws[f'{col_letter}7'].font = Font(bold=True)


def build_volume_assumptions_sheet(wb):
    """Build Volume_Assumptions sheet"""

    print("Building Volume_Assumptions sheet...")
    ws = wb['Volume_Assumptions']

    # Title
    ws['A1'] = 'VOLUME ASSUMPTIONS'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    # Add timeline header
    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # Seasonality block
    ws['A10'] = 'Seasonality Factors'
    ws['A10'].font = Font(bold=True, size=11)

    ws['A11'] = 'Month (1-12)'
    ws['A12'] = 'Month name'
    ws['A13'] = 'Seasonality factor'
    ws['A14'] = 'Check sum'

    # Seasonality values (Southern Hemisphere - Australia/NZ)
    seasonality = [1.20, 1.15, 1.05, 0.95, 0.90, 0.85, 0.85, 0.88, 0.95, 1.00, 1.10, 1.12]
    month_names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

    for i, (factor, name) in enumerate(zip(seasonality, month_names), start=1):
        col_letter = get_column_letter(i + 1)  # Start at column B
        ws[f'{col_letter}11'] = i
        ws[f'{col_letter}12'] = name
        ws[f'{col_letter}13'] = factor
        ws[f'{col_letter}13'].number_format = '0.00'
        ws[f'{col_letter}13'].fill = BLUE_INPUT

    ws['B14'] = '=SUM(B13:M13)'
    ws['B14'].number_format = '0.00'
    ws['B14'].font = Font(bold=True)
    ws['C14'] = '(must = 12.00)'
    ws['C14'].font = Font(italic=True, size=9)

    # Volume assumptions by brand-pack-channel
    ws['A18'] = 'Brand-Pack-Channel'
    ws['A18'].font = Font(bold=True)
    ws['A18'].fill = DARK_BLUE_HEADER
    ws['A18'].font = Font(bold=True, color="FFFFFF")

    ws['B18'] = 'Year 1 Base (HL p.a.)'
    ws['B18'].font = Font(bold=True, color="FFFFFF")
    ws['B18'].fill = DARK_BLUE_HEADER

    ws['C18'] = 'Growth % Base'
    ws['C18'].font = Font(bold=True, color="FFFFFF")
    ws['C18'].fill = DARK_BLUE_HEADER

    ws['D18'] = 'Upside'
    ws['D18'].font = Font(bold=True, color="FFFFFF")
    ws['D18'].fill = DARK_BLUE_HEADER

    ws['E18'] = 'Downside'
    ws['E18'].font = Font(bold=True, color="FFFFFF")
    ws['E18'].fill = DARK_BLUE_HEADER

    # Define all 30 SKUs (5 brands × 3 packs × 2 channels)
    brands = ['SkyBrew', 'AllDark', 'GoldCoast', 'CityLight', 'HarborIPA']
    packs = ['Bottles', 'Cans', 'Kegs']
    channels = ['OnTrade', 'OffTrade']

    # Sample base volumes (illustrative - user can adjust)
    # Total should be ~800,000 HL/year at 80% utilization of 1M HL capacity
    base_volumes = {
        'SkyBrew-Bottles-OnTrade': 50000,
        'SkyBrew-Bottles-OffTrade': 80000,
        'SkyBrew-Cans-OnTrade': 20000,
        'SkyBrew-Cans-OffTrade': 100000,
        'SkyBrew-Kegs-OnTrade': 60000,
        'SkyBrew-Kegs-OffTrade': 0,  # Kegs not sold off-trade

        'AllDark-Bottles-OnTrade': 15000,
        'AllDark-Bottles-OffTrade': 25000,
        'AllDark-Cans-OnTrade': 8000,
        'AllDark-Cans-OffTrade': 30000,
        'AllDark-Kegs-OnTrade': 17000,
        'AllDark-Kegs-OffTrade': 0,

        'GoldCoast-Bottles-OnTrade': 20000,
        'GoldCoast-Bottles-OffTrade': 35000,
        'GoldCoast-Cans-OnTrade': 12000,
        'GoldCoast-Cans-OffTrade': 45000,
        'GoldCoast-Kegs-OnTrade': 25000,
        'GoldCoast-Kegs-OffTrade': 0,

        'CityLight-Bottles-OnTrade': 25000,
        'CityLight-Bottles-OffTrade': 60000,
        'CityLight-Cans-OnTrade': 15000,
        'CityLight-Cans-OffTrade': 70000,
        'CityLight-Kegs-OnTrade': 10000,
        'CityLight-Kegs-OffTrade': 0,

        'HarborIPA-Bottles-OnTrade': 18000,
        'HarborIPA-Bottles-OffTrade': 22000,
        'HarborIPA-Cans-OnTrade': 10000,
        'HarborIPA-Cans-OffTrade': 25000,
        'HarborIPA-Kegs-OnTrade': 15000,
        'HarborIPA-Kegs-OffTrade': 0,
    }

    # Growth rates by brand type
    growth_rates = {
        'SkyBrew': (0.03, 0.05, 0.01),  # Core brand
        'AllDark': (0.04, 0.06, 0.02),  # Craft
        'GoldCoast': (0.04, 0.06, 0.02),  # Craft
        'CityLight': (0.025, 0.04, 0.01),  # Core but slower
        'HarborIPA': (0.05, 0.07, 0.03),  # Premium, fastest growth
    }

    row = 19
    for brand in brands:
        for pack in packs:
            for channel in channels:
                sku = f'{brand}-{pack}-{channel}'
                base_vol = base_volumes.get(sku, 0)
                base_growth, upside_growth, downside_growth = growth_rates[brand]

                ws[f'A{row}'] = sku
                ws[f'B{row}'] = base_vol
                ws[f'B{row}'].number_format = '#,##0'
                ws[f'B{row}'].fill = BLUE_INPUT

                ws[f'C{row}'] = base_growth
                ws[f'C{row}'].number_format = '0.0%'
                ws[f'C{row}'].fill = BLUE_INPUT

                ws[f'D{row}'] = upside_growth
                ws[f'D{row}'].number_format = '0.0%'
                ws[f'D{row}'].fill = BLUE_INPUT

                ws[f'E{row}'] = downside_growth
                ws[f'E{row}'].number_format = '0.0%'
                ws[f'E{row}'].fill = BLUE_INPUT

                # Monthly volume formulas (columns F onwards)
                for month in range(1, MONTHS + 1):
                    col = TIMELINE_START_COL + month - 1
                    col_letter = get_column_letter(col)
                    year_cell = f'{col_letter}5'  # Year decimal from header

                    # Formula: (Base/12) * (1+Growth)^Year * Seasonality
                    # Seasonality lookup: INDEX($B$13:$M$13, MOD(month-1, 12)+1)
                    formula = f'=($B{row}/12)*(1+INDEX($C{row}:$E{row},1,Control!$B$10))^{year_cell}*INDEX($B$13:$M$13,1,MOD({month}-1,12)+1)'
                    ws[f'{col_letter}{row}'] = formula
                    ws[f'{col_letter}{row}'].number_format = '#,##0'

                row += 1

    # Summary totals
    total_row = row + 1
    ws[f'A{total_row}'] = 'TOTAL HL SOLD'
    ws[f'A{total_row}'].font = Font(bold=True)
    ws[f'A{total_row}'].fill = YELLOW_OUTPUT

    # Sum formulas for each month
    for month in range(1, MONTHS + 1):
        col = TIMELINE_START_COL + month - 1
        col_letter = get_column_letter(col)
        ws[f'{col_letter}{total_row}'] = f'=SUM({col_letter}19:{col_letter}{row-1})'
        ws[f'{col_letter}{total_row}'].number_format = '#,##0'
        ws[f'{col_letter}{total_row}'].fill = YELLOW_OUTPUT
        ws[f'{col_letter}{total_row}'].font = Font(bold=True)

    # Set column widths
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12

    for month in range(1, min(MONTHS + 1, 25)):  # Set width for first 24 months
        col = TIMELINE_START_COL + month - 1
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 12

    print(f"  ✓ Volume_Assumptions sheet complete (30 SKUs × {MONTHS} months)")


def build_pricing_revenue_sheet(wb):
    """Build Pricing_Revenue sheet"""

    print("Building Pricing_Revenue sheet...")
    ws = wb['Pricing_Revenue']

    # Title
    ws['A1'] = 'PRICING & REVENUE'
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = DARK_BLUE_HEADER
    ws.merge_cells('A1:E1')

    # Add timeline header
    add_timeline_header(ws, start_col=TIMELINE_START_COL, num_months=MONTHS)

    # Gross price assumptions
    ws['A10'] = 'Gross Price ($/HL)'
    ws['A10'].font = Font(bold=True, size=11)

    headers = ['Brand-Pack-Channel', 'Year 1 $/HL', 'Esc % Base', 'Upside', 'Downside']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=11, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = DARK_BLUE_HEADER

    # Sample gross prices ($/HL) - user can adjust
    # OnTrade typically 10-20% higher than OffTrade
    gross_prices = {
        'SkyBrew-Bottles-OnTrade': 180,
        'SkyBrew-Bottles-OffTrade': 160,
        'SkyBrew-Cans-OnTrade': 175,
        'SkyBrew-Cans-OffTrade': 155,
        'SkyBrew-Kegs-OnTrade': 170,

        'AllDark-Bottles-OnTrade': 200,
        'AllDark-Bottles-OffTrade': 180,
        'AllDark-Cans-OnTrade': 195,
        'AllDark-Cans-OffTrade': 175,
        'AllDark-Kegs-OnTrade': 190,

        'GoldCoast-Bottles-OnTrade': 210,
        'GoldCoast-Bottles-OffTrade': 190,
        'GoldCoast-Cans-OnTrade': 205,
        'GoldCoast-Cans-OffTrade': 185,
        'GoldCoast-Kegs-OnTrade': 200,

        'CityLight-Bottles-OnTrade': 160,
        'CityLight-Bottles-OffTrade': 140,
        'CityLight-Cans-OnTrade': 155,
        'CityLight-Cans-OffTrade': 135,
        'CityLight-Kegs-OnTrade': 150,

        'HarborIPA-Bottles-OnTrade': 240,
        'HarborIPA-Bottles-OffTrade': 220,
        'HarborIPA-Cans-OnTrade': 235,
        'HarborIPA-Cans-OffTrade': 215,
        'HarborIPA-Kegs-OnTrade': 230,
    }

    # Price escalation (same for all SKUs in Base, but scenario-dependent)
    price_esc = (0.025, 0.035, 0.015)  # Base, Upside, Downside

    brands = ['SkyBrew', 'AllDark', 'GoldCoast', 'CityLight', 'HarborIPA']
    packs = ['Bottles', 'Cans', 'Kegs']
    channels = ['OnTrade', 'OffTrade']

    row = 12
    for brand in brands:
        for pack in packs:
            for channel in channels:
                sku = f'{brand}-{pack}-{channel}'
                price = gross_prices.get(sku, 0)

                if price == 0:
                    continue  # Skip if no price (e.g., Kegs-OffTrade)

                ws[f'A{row}'] = sku
                ws[f'B{row}'] = price
                ws[f'B{row}'].number_format = '$#,##0'
                ws[f'B{row}'].fill = BLUE_INPUT

                ws[f'C{row}'] = price_esc[0]
                ws[f'C{row}'].number_format = '0.0%'
                ws[f'C{row}'].fill = BLUE_INPUT

                ws[f'D{row}'] = price_esc[1]
                ws[f'D{row}'].number_format = '0.0%'
                ws[f'D{row}'].fill = BLUE_INPUT

                ws[f'E{row}'] = price_esc[2]
                ws[f'E{row}'].number_format = '0.0%'
                ws[f'E{row}'].fill = BLUE_INPUT

                row += 1

    # Discount & promo assumptions
    discount_row = row + 2
    ws[f'A{discount_row}'] = 'Trade Discounts & Promos'
    ws[f'A{discount_row}'].font = Font(bold=True, size=11)

    ws[f'A{discount_row+1}'] = 'Channel'
    ws[f'B{discount_row+1}'] = 'Trade Discount %'
    ws[f'C{discount_row+1}'] = 'Promo %'
    ws[f'D{discount_row+1}'] = 'Total Deduction %'

    for col in ['A', 'B', 'C', 'D']:
        ws[f'{col}{discount_row+1}'].font = Font(bold=True, color="FFFFFF")
        ws[f'{col}{discount_row+1}'].fill = DARK_BLUE_HEADER

    ws[f'A{discount_row+2}'] = 'OnTrade'
    ws[f'B{discount_row+2}'] = 0.15
    ws[f'B{discount_row+2}'].number_format = '0%'
    ws[f'B{discount_row+2}'].fill = BLUE_INPUT

    ws[f'C{discount_row+2}'] = 0.05
    ws[f'C{discount_row+2}'].number_format = '0%'
    ws[f'C{discount_row+2}'].fill = BLUE_INPUT

    ws[f'D{discount_row+2}'] = f'=B{discount_row+2}+C{discount_row+2}'
    ws[f'D{discount_row+2}'].number_format = '0%'

    ws[f'A{discount_row+3}'] = 'OffTrade'
    ws[f'B{discount_row+3}'] = 0.20
    ws[f'B{discount_row+3}'].number_format = '0%'
    ws[f'B{discount_row+3}'].fill = BLUE_INPUT

    ws[f'C{discount_row+3}'] = 0.06
    ws[f'C{discount_row+3}'].number_format = '0%'
    ws[f'C{discount_row+3}'].fill = BLUE_INPUT

    ws[f'D{discount_row+3}'] = f'=B{discount_row+3}+C{discount_row+3}'
    ws[f'D{discount_row+3}'].number_format = '0%'

    # Set column widths
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 18

    print(f"  ✓ Pricing_Revenue sheet complete")
    print("    (Note: Net revenue and excise calculations will be added in later version)")


def main():
    """Main execution"""
    print("\n" + "="*60)
    print("Building Volume & Pricing Sheets")
    print("="*60 + "\n")

    # Load workbook
    wb = load_workbook('Brewery_Financial_Model.xlsx')

    # Build sheets
    build_volume_assumptions_sheet(wb)
    build_pricing_revenue_sheet(wb)

    # Save
    wb.save('Brewery_Financial_Model.xlsx')

    print("\n✅ Volume & Pricing sheets complete!")
    print(f"\n   - {MONTHS} months of projections created")
    print("   - 30 SKUs (5 brands × 3 packs × 2 channels)")
    print("   - Seasonality applied")
    print("   - Scenario-driven growth rates")
    print("\nNext step: Run 03_operations_costs.py")


if __name__ == '__main__':
    main()
